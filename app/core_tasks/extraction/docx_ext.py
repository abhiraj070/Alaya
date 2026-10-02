from docx import Document

def extract_text(file_path: str) -> str:
    document = Document(file_path)

    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                cells.append(cell.text.strip())

            parts.append(" | ".join(cells))

    return "\n".join(parts)