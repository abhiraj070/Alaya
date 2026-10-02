from pptx import Presentation

def extract_text(file_path: str) -> str:
    presentation = Presentation(file_path)
    slides = []
    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_parts = []
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = shape.text.strip()
            if text:
                slide_parts.append(text)
        if slide_parts:
            slides.append(
                f"[Slide {slide_number}]\n"
                + "\n".join(slide_parts)
            )
    return "\n\n".join(slides)