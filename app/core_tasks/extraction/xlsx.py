from openpyxl import load_workbook

def extract_text(file_path: str) -> str:
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]
                if values:
                    rows.append(" | ".join(values))
            if rows:
                sheet_text = (
                    f"[Sheet: {worksheet.title}]\n"
                    + "\n".join(rows)
                )
                sheets.append(sheet_text)
    finally:
        workbook.close()
    return "\n\n".join(sheets)